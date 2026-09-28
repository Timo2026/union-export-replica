#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
CAD Quote - 画图+报价组合 SKILL

用法:
    python3 cad_quote.py [command] [args]
    
示例:
    python3 cad_quote.py flange 60 50 15 20
    python3 cad_quote.py box 80 60 40
    python3 cad_quote.py cyl 25 50
"""

import sys
import os
import json
import sqlite3
from datetime import datetime
from pathlib import Path

# 添加 SKILL 路径
SKILL_DIR = Path(__file__).parent.parent
CAD_SCRIPT_DIR = SKILL_DIR.parent / "step-factory/scripts"
sys.path.insert(0, str(CAD_SCRIPT_DIR))

HOME = Path.home()
DATE = datetime.now().strftime("%Y-%m-%d")
DATETIME = datetime.now().strftime("%Y%m%d_%H%M%S")
CAD_DIR = HOME / f"STEP/{DATE}/cad_files"
QUOTES_DIR = HOME / f"STEP/{DATE}/quotes"
CAD_RAG_DIR = HOME / ".openclaw/cad_rag"

os.makedirs(CAD_DIR, exist_ok=True)
os.makedirs(QUOTES_DIR, exist_ok=True)

# 导入 CAD 模块
from step_occ import (
    make_box, make_cylinder, make_flange, make_tube,
    make_cone, make_sphere, make_torus, make_shaft,
    make_flange_with_holes
)
from OCC.Core.STEPControl import STEPControl_Writer

# 导入报价模块
import openpyxl
from openpyxl.styles import Font, Alignment


def get_quote_from_rag(part_type, material="304不锈钢"):
    """从 RAG 知识库查询报价"""
    db_path = CAD_RAG_DIR / "metadata.db"
    
    if not db_path.exists():
        return None
    
    conn = sqlite3.connect(db_path)
    cursor = conn.cursor()
    
    # 模糊匹配
    cursor.execute("""
        SELECT part_name, material, quantity, unit_price 
        FROM quotes 
        WHERE part_name LIKE ? OR material LIKE ?
        LIMIT 1
    """, (f"%{part_type}%", f"%{material}%"))
    
    result = cursor.fetchone()
    conn.close()
    
    if result:
        return {"part_name": result[0], "material": result[1], "quantity": result[2], "unit_price": result[3]}
    return None


def export_step(shape, filename):
    """导出 STEP 文件"""
    path = CAD_DIR / filename
    writer = STEPControl_Writer()
    writer.Transfer(shape, 1)
    writer.Write(str(path))
    return path


def generate_quote_xlsx(part_name, material, quantity, unit_price, step_filename):
    """生成报价单 XLSX"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "报价单"
    
    # 列宽
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 15
    
    # 标题
    ws['A1'] = "CNC加工报价单"
    ws['A1'].font = Font(size=16, bold=True)
    ws['A1'].alignment = Alignment(horizontal='center')
    ws.merge_cells('A1:E1')
    
    # 信息
    ws['A3'] = "报价单号:"
    ws['B3'] = f"QUOTE-{DATETIME}"
    ws['A4'] = "日期:"
    ws['B4'] = DATE
    ws['A5'] = "3D图纸:"
    ws['B5'] = step_filename
    
    # 表头
    headers = ["零件名称", "材料", "数量", "单价(元)", "总价(元)"]
    for i, h in enumerate(headers, 1):
        cell = ws.cell(row=7, column=i, value=h)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center')
    
    # 数据
    total = quantity * unit_price
    data = [
        [part_name, material, quantity, unit_price, total],
    ]
    
    for row_idx, row_data in enumerate(data, 8):
        for col_idx, value in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    
    # 保存
    xlsx_path = QUOTES_DIR / f"报价单_{DATETIME}.xlsx"
    wb.save(str(xlsx_path))
    
    return xlsx_path


def cmd_box(w, h, d):
    """方块"""
    shape = make_box(w, h, d)
    filename = f"box_{w}x{h}x{d}_{DATETIME}.step"
    path = export_step(shape, filename)
    
    # 查报价
    quote = get_quote_from_rag("box") or {"part_name": f"方块 {w}x{h}x{d}", "material": "碳钢", "quantity": 1, "unit_price": 200.0}
    
    # 生成报价单
    xlsx_path = generate_quote_xlsx(quote["part_name"], quote["material"], quote["quantity"], quote["unit_price"], filename)
    
    return path, xlsx_path


def cmd_cylinder(r, h):
    """圆柱"""
    shape = make_cylinder(r, h)
    filename = f"cyl_r{r}h{h}_{DATETIME}.step"
    path = export_step(shape, filename)
    
    # 查报价
    quote = get_quote_from_rag("圆柱") or {"part_name": f"圆柱 r{r}h{h}", "material": "304不锈钢", "quantity": 10, "unit_price": 50.0}
    
    # 生成报价单
    xlsx_path = generate_quote_xlsx(quote["part_name"], quote["material"], quote["quantity"], quote["unit_price"], filename)
    
    return path, xlsx_path


def cmd_flange(outer_r, inner_r, hole_r, thickness):
    """法兰"""
    shape = make_flange(outer_r, inner_r, hole_r, thickness)
    filename = f"flange_{outer_r}_{inner_r}_{hole_r}_{thickness}_{DATETIME}.step"
    path = export_step(shape, filename)
    
    # 查报价
    quote = get_quote_from_rag("法兰") or {"part_name": f"法兰盘 {outer_r}x{inner_r}x{hole_r}", "material": "304不锈钢", "quantity": 5, "unit_price": 580.0}
    
    # 生成报价单
    xlsx_path = generate_quote_xlsx(quote["part_name"], quote["material"], quote["quantity"], quote["unit_price"], filename)
    
    return path, xlsx_path


def main():
    print("=" * 60)
    print("CAD Quote - 画图+报价组合")
    print("=" * 60)
    
    if len(sys.argv) < 2:
        print("\n用法:")
        print("  python3 cad_quote.py box W H D")
        print("  python3 cad_quote.py cyl r h")
        print("  python3 cad_quote.py flange outer_r inner_r hole_r thickness")
        print("\n示例:")
        print("  python3 cad_quote.py box 80 60 40")
        print("  python3 cad_quote.py cyl 25 50")
        print("  python3 cad_quote.py flange 60 50 15 20")
        return
    
    cmd = sys.argv[1]
    
    try:
        if cmd == "box" and len(sys.argv) >= 5:
            w, h, d = float(sys.argv[2]), float(sys.argv[3]), float(sys.argv[4])
            step_path, xlsx_path = cmd_box(w, h, d)
            
        elif cmd == "cyl" and len(sys.argv) >= 4:
            r, h = float(sys.argv[2]), float(sys.argv[3])
            step_path, xlsx_path = cmd_cylinder(r, h)
            
        elif cmd == "flange" and len(sys.argv) >= 6:
            outer_r, inner_r, hole_r, thickness = map(float, sys.argv[2:6])
            step_path, xlsx_path = cmd_flange(outer_r, inner_r, hole_r, thickness)
            
        else:
            print(f"❌ 未知命令或参数不足: {cmd}")
            return
        
        print(f"\n✅ 生成完成!")
        print(f"   STEP: {step_path.name}")
        print(f"   报价: {xlsx_path.name}")
        print(f"\n📁 输出目录: /home/<user>/STEP/{DATE}/")
        
    except Exception as e:
        print(f"\n❌ 错误: {e}")
        import traceback
        traceback.print_exc()


if __name__ == "__main__":
    main()